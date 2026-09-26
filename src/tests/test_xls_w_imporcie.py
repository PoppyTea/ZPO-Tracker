"""
`.xls` w imporcie miesięcy i w porządkowaniu (ZPO-68).

Ten sam miesiąc zapisany jako `.xlsx` i jako `.xls` ma dawać to samo
znaczenie. Porównujemy wiersze PO WALIDACJI (`WierszImportu`), bo tam
widać realną różnicę: data z `.xls` jest liczbą zmiennoprzecinkową,
a pusta komórka to `''`, nie `None`.
"""
import datetime as dt

import openpyxl
import pytest
import xlwt

from zpo_tracker import import_orchestrator, porzadkowanie
from zpo_tracker.gui.zakladka_import_export import _wczytaj_surowe_wiersze

NAGLOWKI = list(import_orchestrator.MAPA_NAGLOWKOW)

# data, nadawca, adres, kurier, rejon, total, zpo, pni, vinted, automaty,
# k48, niezrealizowane, wykonawca — kolejność jak w MAPA_NAGLOWKOW
WIERSZE = [
    [dt.datetime(2026, 8, 3), "Żabka", "Kwiatowa 8", "Kowalski Jan", "WA87",
     12, 3, "228648", None, None, None, None, "Firma A"],
    [dt.datetime(2026, 8, 3), "Rossmann", "Polna 12/14", "Kowalski Jan", None,
     7, None, None, None, 2, None, None, "Firma A"],
    # wiersz-szablon bez daty i kuriera W ŚRODKU arkusza (na końcu .xls
    # w ogóle by go nie zapisał): ma zostać pominięty w obu formatach
    [None, None, None, None, None, None, None, None, None, None, None, None, None],
    [dt.datetime(2026, 8, 4), "ZUS", "Marsa 56 m. 3", "Nowak Anna", "WA101",
     1, None, None, None, None, None, None, "Firma B"],
]

# porządkowanie: bliskie powtórzenie (scalić, zsumować) i dalekie (zostawić)
NAGLOWKI_PORZ = ["data", "Kurier", " Pełna Nazwa Nadawcy",
                 "Adres odbioru dla wszystkich nadawców",
                 " Wpisujemy łączną liczbę odebranych Pocztexów", "PNI ZPO"]


def _wiersze_porz():
    d = dt.datetime(2026, 8, 3)
    wiersze = [[d, "Kowalski Jan", "Żabka", "Kwiatowa 8", 2, None],
               [d, "Kowalski Jan", "Żabka", "Kwiatowa 8", 4, "228648"]]
    wiersze += [[d, f"Kurier {i}", "Inny", f"Ulica {i}", 1, None] for i in range(15)]
    wiersze += [[d, "Kowalski Jan", "Rossmann", "Polna 1", 5, None]]
    wiersze += [[dt.datetime(2026, 8, 4), "Kurier X", "Rossmann", "Polna 1", 3, None]]
    wiersze += [[d, "Kowalski Jan", "Rossmann", "Polna 1", 6, None]]
    return wiersze


def _zapisz_xlsx(sciezka, naglowki, wiersze):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sierpień"
    ws.append(naglowki)
    for w in wiersze:
        ws.append(w)
    wb.save(sciezka)
    return sciezka


def _zapisz_xls(sciezka, naglowki, wiersze):
    """Tak, jak zapisuje stary Excel: data jako liczba z formatem daty,
    pusta komórka po prostu nieobecna."""
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Sierpień")
    styl_daty = xlwt.easyxf(num_format_str="YYYY-MM-DD")
    for k, n in enumerate(naglowki):
        ws.write(0, k, n)
    for r, w in enumerate(wiersze, start=1):
        for k, v in enumerate(w):
            if v is None:
                continue
            if isinstance(v, dt.datetime):
                ws.write(r, k, v, styl_daty)
            else:
                ws.write(r, k, v)
    wb.save(str(sciezka))
    return sciezka


def _znaczenie(zwalidowane):
    return [w.model_dump(exclude={"numer_wiersza"}) for w in zwalidowane]


@pytest.fixture
def miesiac(tmp_path):
    return (_zapisz_xlsx(tmp_path / "sierpien.xlsx", NAGLOWKI, WIERSZE),
            _zapisz_xls(tmp_path / "sierpien.xls", NAGLOWKI, WIERSZE))


def test_xls_wczytuje_sie_w_imporcie_miesiaca(miesiac):
    _, xls = miesiac
    surowe = _wczytaj_surowe_wiersze(xls)
    assert len(surowe) == len(WIERSZE)


def test_xls_daje_po_walidacji_to_samo_co_xlsx(miesiac):
    xlsx, xls = miesiac
    ok_xlsx, odrz_xlsx = import_orchestrator.zwaliduj_wiersze(_wczytaj_surowe_wiersze(xlsx))
    ok_xls, odrz_xls = import_orchestrator.zwaliduj_wiersze(_wczytaj_surowe_wiersze(xls))
    assert len(ok_xlsx) == 3 and not odrz_xlsx
    assert not odrz_xls, [o["powod"] for o in odrz_xls]
    assert _znaczenie(ok_xls) == _znaczenie(ok_xlsx)


def test_xls_zachowuje_numery_wierszy_zrodla(miesiac):
    xlsx, xls = miesiac
    klucz = import_orchestrator.KLUCZ_NUMERU_WIERSZA
    numery = lambda p: [w.get(klucz) for w in _wczytaj_surowe_wiersze(p)]
    assert numery(xls) == numery(xlsx)


def test_xls_z_rozszerzeniem_xlsx_tez_sie_wczytuje(tmp_path):
    """Format po zawartości, nie po nazwie: pliki przychodzą przez
    przeglądarkę i bywają przemianowane."""
    podszyty = _zapisz_xls(tmp_path / "przemianowany.xlsx", NAGLOWKI, WIERSZE)
    ok, odrz = import_orchestrator.zwaliduj_wiersze(_wczytaj_surowe_wiersze(podszyty))
    assert len(ok) == 3 and not odrz


def _wynik(sciezka):
    ws = openpyxl.load_workbook(sciezka).active
    return [list(r) for r in ws.iter_rows(values_only=True)]


def test_porzadkowanie_dziala_na_xls_tak_samo_jak_na_xlsx(tmp_path):
    wiersze = _wiersze_porz()
    xlsx = _zapisz_xlsx(tmp_path / "p.xlsx", NAGLOWKI_PORZ, wiersze)
    xls = _zapisz_xls(tmp_path / "p.xls", NAGLOWKI_PORZ, wiersze)
    r_xlsx = porzadkowanie.porzadkuj(xlsx, tmp_path / "wynik_xlsx.xlsx")
    r_xls = porzadkowanie.porzadkuj(xls, tmp_path / "wynik_xls.xlsx")

    assert r_xls.sumy_zgodne
    assert (r_xls.wierszy_przed, r_xls.wierszy_po, r_xls.scalone, r_xls.do_decyzji) == \
           (r_xlsx.wierszy_przed, r_xlsx.wierszy_po, r_xlsx.scalone, r_xlsx.do_decyzji)


def test_porzadkowanie_xls_nie_zamienia_dat_na_liczby(tmp_path):
    wiersze = _wiersze_porz()
    xlsx = _zapisz_xlsx(tmp_path / "p.xlsx", NAGLOWKI_PORZ, wiersze)
    xls = _zapisz_xls(tmp_path / "p.xls", NAGLOWKI_PORZ, wiersze)
    porzadkowanie.porzadkuj(xlsx, tmp_path / "wynik_xlsx.xlsx")
    porzadkowanie.porzadkuj(xls, tmp_path / "wynik_xls.xlsx")

    daty = lambda p: [w[0] for w in _wynik(p)[1:]]
    assert daty(tmp_path / "wynik_xls.xlsx") == daty(tmp_path / "wynik_xlsx.xlsx")


# --- okno wyboru pliku ----------------------------------------------------

def _ma_display():
    import subprocess
    import sys
    try:
        wynik = subprocess.run(
            [sys.executable, "-c",
             "import tkinter as tk; r = tk.Tk(); tk.Entry(r).pack(); r.update(); r.destroy()"],
            capture_output=True, timeout=10,
        )
        return wynik.returncode == 0
    except Exception:
        return False


@pytest.mark.skipif(not _ma_display(), reason="wymaga środowiska graficznego (DISPLAY)")
def test_okno_importu_pokazuje_pliki_xls(tmp_path, monkeypatch):
    """Czytnik obsługuje .xls, ale bez filtra w oknie wyboru użytkownik nie
    ma jak takiego pliku wskazać — import rejonarza obok ma oba formaty."""
    import tkinter as tk

    from zpo_tracker import repo
    from zpo_tracker.gui import zakladka_import_export as zakladka_mod

    podane = {}

    def udawane_okno(**kwargs):
        podane.update(kwargs)
        return ""

    monkeypatch.setattr(zakladka_mod.filedialog, "askopenfilename", udawane_okno)
    root = tk.Tk()
    conn = repo.polacz(":memory:")
    repo.utworz_schemat(conn)
    try:
        zakladka_mod.ZakladkaImportExport(root, conn, tmp_path).importuj()
    finally:
        conn.close()
        root.destroy()

    wzorce = " ".join(w for _, w in podane["filetypes"])
    assert "*.xls" in wzorce.split() and "*.xlsx" in wzorce.split()


@pytest.mark.skipif(not _ma_display(), reason="wymaga środowiska graficznego (DISPLAY)")
def test_plik_ktory_nie_jest_arkuszem_daje_komunikat_nie_wyjatek(tmp_path, monkeypatch):
    """„Wszystkie pliki” pozwala wskazać np. stronę logowania zapisaną przez
    przeglądarkę. Nieobsłużony wyjątek z callbacku Tk jest w buildzie
    `console=False` niewidoczny — użytkownik ma dostać komunikat, jak przy
    imporcie rejonarza."""
    import tkinter as tk

    from zpo_tracker import repo
    from zpo_tracker.gui import zakladka_import_export as zakladka_mod

    nie_arkusz = tmp_path / "logowanie.xlsx"
    nie_arkusz.write_text("<html>zaloguj się</html>", encoding="utf-8")
    bledy = []
    monkeypatch.setattr(zakladka_mod.filedialog, "askopenfilename", lambda **_: str(nie_arkusz))
    monkeypatch.setattr(zakladka_mod.messagebox, "showerror",
                        lambda tytul, tresc, **_: bledy.append((tytul, tresc)))
    root = tk.Tk()
    conn = repo.polacz(":memory:")
    repo.utworz_schemat(conn)
    try:
        zakladka = zakladka_mod.ZakladkaImportExport(root, conn, tmp_path)
        zakladka.importuj()
        etykieta = zakladka.etykieta_import.cget("text")
    finally:
        conn.close()
        root.destroy()

    assert len(bledy) == 1 and "logowanie.xlsx" in bledy[0][1]
    assert "Nie wczytano" in etykieta
